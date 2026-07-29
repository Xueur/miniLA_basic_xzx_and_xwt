import csv
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUT = ROOT.parent / "实验一_单周期CPU完整设计说明.docx"

BLUE = "1F4E78"
LIGHT_BLUE = "DDEBF7"
LIGHT_ORANGE = "FCE4D6"
LIGHT_GREEN = "E2F0D9"
GRID = "9EADBA"
MUTED = RGBColor(90, 100, 110)

sample = ["SLLI", "ADDI", "LUI", "ORI", "LW", "JAL", "BEQ", "BNE"]
group_a = ["SLL", "SRL", "SRLI", "SRA", "SRAI", "ADD", "SUB", "AUIPC", "XOR", "XORI", "LB", "LBU", "LH", "LHU", "SW", "SB", "SH", "JALR"]
group_b = ["MUL", "MULH", "MULHU", "DIV", "DIVU", "REM", "REMU", "SLT", "SLTI", "SLTU", "SLTIU", "OR", "AND", "ANDI", "BLT", "BGE", "BLTU", "BGEU"]

types = {
    **{x: "R" for x in ["SLL", "SRL", "SRA", "ADD", "SUB", "XOR", "MUL", "MULH", "MULHU", "DIV", "DIVU", "REM", "REMU", "SLT", "SLTU", "OR", "AND"]},
    **{x: "I" for x in ["SLLI", "SRLI", "SRAI", "ADDI", "ORI", "XORI", "LB", "LBU", "LH", "LHU", "LW", "JALR", "SLTI", "SLTIU", "ANDI"]},
    **{x: "S" for x in ["SW", "SB", "SH"]},
    **{x: "B" for x in ["BEQ", "BNE", "BLT", "BGE", "BLTU", "BGEU"]},
    "LUI": "U", "AUIPC": "U", "JAL": "J",
}

f3 = {
    "ADDI": "000", "SLLI": "001", "SLTI": "010", "SLTIU": "011", "XORI": "100", "SRLI": "101", "SRAI": "101", "ORI": "110", "ANDI": "111",
    "ADD": "000", "SUB": "000", "SLL": "001", "SLT": "010", "SLTU": "011", "XOR": "100", "SRL": "101", "SRA": "101", "OR": "110", "AND": "111",
    "MUL": "000", "MULH": "001", "MULHU": "011", "DIV": "100", "DIVU": "101", "REM": "110", "REMU": "111",
    "LB": "000", "LH": "001", "LW": "010", "LBU": "100", "LHU": "101",
    "SB": "000", "SH": "001", "SW": "010",
    "BEQ": "000", "BNE": "001", "BLT": "100", "BGE": "101", "BLTU": "110", "BGEU": "111", "JALR": "000",
}

f7 = {x: "0000000" for x in ["SLLI", "SRLI", "ADD", "SLL", "SLT", "SLTU", "XOR", "SRL", "OR", "AND"]}
f7.update({"SRAI": "0100000", "SUB": "0100000", "SRA": "0100000"})
f7.update({x: "0000001" for x in ["MUL", "MULH", "MULHU", "DIV", "DIVU", "REM", "REMU"]})

alu = {
    "ADDI": "ADD", "ADD": "ADD", "AUIPC": "ADD", "LB": "ADD", "LBU": "ADD", "LH": "ADD", "LHU": "ADD", "LW": "ADD", "SB": "ADD", "SH": "ADD", "SW": "ADD", "JALR": "ADD",
    "SUB": "SUB", "XOR": "XOR", "XORI": "XOR", "OR": "OR", "ORI": "OR", "AND": "AND", "ANDI": "AND",
    "SLL": "SLL", "SLLI": "SLL", "SRL": "SRL", "SRLI": "SRL", "SRA": "SRA", "SRAI": "SRA",
    "SLT": "SLT", "SLTI": "SLT", "BLT": "SLT", "BGE": "SGE", "SLTU": "SLTU", "SLTIU": "SLTU", "BLTU": "SLTU", "BGEU": "SGEU",
    "BEQ": "EQ", "BNE": "NE", "MUL": "MUL", "MULH": "MULH", "MULHU": "MULHU", "DIV": "DIV", "DIVU": "DIVU", "REM": "REM", "REMU": "REMU",
}


def opcode(inst):
    if inst in ["LUI"]: return "0110111"
    if inst in ["AUIPC"]: return "0010111"
    if inst in ["JAL"]: return "1101111"
    if inst in ["JALR"]: return "1100111"
    if inst in ["LB", "LBU", "LH", "LHU", "LW"]: return "0000011"
    if inst in ["SB", "SH", "SW"]: return "0100011"
    if inst in ["BEQ", "BNE", "BLT", "BGE", "BLTU", "BGEU"]: return "1100011"
    if types[inst] == "I": return "0010011"
    return "0110011"


def group(inst):
    return "示例" if inst in sample else "A" if inst in group_a else "B"


def control(inst):
    t = types[inst]
    is_load = inst in ["LB", "LBU", "LH", "LHU", "LW"]
    is_store = inst in ["SB", "SH", "SW"]
    is_branch = t == "B"
    npc = "JALR" if inst == "JALR" else "JMP" if inst == "JAL" else "BRA" if is_branch else "PC4"
    rf_we = "0" if is_store or is_branch else "1"
    wb = "-" if rf_we == "0" else "RAM" if is_load else "PC4" if inst in ["JAL", "JALR"] else "EXT" if inst == "LUI" else "ALU"
    ext = "S" if is_store else "B" if is_branch else "U" if inst in ["LUI", "AUIPC"] else "J" if inst == "JAL" else "-" if t == "R" else "I"
    if inst in ["LUI", "JAL"]:
        aa, ab = "-", "-"
    else:
        aa = "PC" if inst == "AUIPC" else "RS1"
        ab = "EXT" if t == "I" or is_store or inst == "AUIPC" else "RS2"
    rop = {"LB": "B", "LBU": "BU", "LH": "H", "LHU": "HU", "LW": "W"}.get(inst, "N")
    wop = {"SB": "B", "SH": "H", "SW": "W"}.get(inst, "N")
    return [npc, rf_we, wb, ext, aa, ab, alu.get(inst, "-"), rop, wop,
            "1" if inst in ["MUL", "MULH", "MULHU"] else "0",
            "1" if inst in ["DIV", "DIVU", "REM", "REMU"] else "0"]


all_insts = sample + group_a + group_b


def rows():
    out = []
    for inst in all_insts:
        out.append([group(inst), inst, types[inst], opcode(inst), f3.get(inst, "-"), f7.get(inst, "-")] + control(inst))
    return out


def datapath_row(inst):
    c = control(inst)
    t = types[inst]
    src_a = "PC" if inst == "AUIPC" else "-" if inst in ["LUI", "JAL"] else "RF[rs1]"
    src_b = "RF[rs2]" if t in ["R", "B"] else "-" if inst in ["LUI", "JAL"] else "SEXT.ext"
    mem = {"LB": "读字节/符号扩展", "LBU": "读字节/零扩展", "LH": "读半字/符号扩展", "LHU": "读半字/零扩展", "LW": "读字", "SB": "写字节", "SH": "写半字", "SW": "写字"}.get(inst, "-")
    wb = {"RAM": "MEXT.ext", "PC4": "NPC.pc4", "EXT": "SEXT.ext", "ALU": "ALU.c", "-": "-"}[c[2]]
    npc = "(rs1+imm)&~1" if inst == "JALR" else "PC+imm" if inst == "JAL" else "条件成立PC+imm，否则PC+4" if t == "B" else "PC+4"
    return [group(inst), inst, types[inst], src_a, src_b, c[6], mem, wb, npc]


def write_csvs():
    with (DOCS / "miniRV_44条指令控制信号表.csv").open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["分组", "指令", "类型", "opcode", "funct3", "funct7", "npc_op", "rf_we", "rf_wsel", "sext_op", "alua_sel", "alub_sel", "alu_op", "ram_rop", "ram_wop", "is_mul", "is_div"])
        w.writerows(rows())
    with (DOCS / "miniRV_44条指令数据通路表.csv").open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["分组", "指令", "类型", "ALU输入A", "ALU输入B", "ALU操作", "访存操作", "写回来源", "下一PC"])
        w.writerows(datapath_row(x) for x in all_insts)


def font(size, bold=False):
    path = Path(__file__).resolve().parents[3] / "tmp/fonts" / ("SourceHanSansSC-Bold.otf" if bold else "SourceHanSansSC-Regular.otf")
    return ImageFont.truetype(str(path), size)


def draw_diagram():
    im = Image.new("RGB", (2400, 1500), "white")
    d = ImageDraw.Draw(im)
    title = font(54, True)
    h = font(30, True)
    sm = font(23)
    tiny = font(18)
    d.text((80, 50), "miniRV 单周期 CPU 完整数据通路", font=title, fill="#17365D")
    d.text((82, 118), "支持 44 条指令；乘除法与访存通过 busy/valid 握手完成", font=sm, fill="#5A6670")

    boxes = {
        "Controller": (850, 190, 1250, 390, "#DDEBF7"),
        "Inst_ROM": (90, 420, 390, 580, "#E2F0D9"),
        "PC": (100, 770, 350, 930, "#E2F0D9"),
        "NPC": (120, 1110, 430, 1290, "#E2F0D9"),
        "RF": (600, 520, 900, 760, "#FFF2CC"),
        "SEXT": (600, 900, 900, 1080, "#FFF2CC"),
        "ALU / MULDIV": (1120, 520, 1500, 760, "#FCE4D6"),
        "MREQ": (1630, 520, 1920, 720, "#FCE4D6"),
        "Data_RAM": (2050, 520, 2350, 720, "#E2F0D9"),
        "MEXT": (2050, 900, 2350, 1080, "#FCE4D6"),
        "WB MUX": (1260, 1120, 1570, 1300, "#D9EAD3"),
    }
    for name, (x1, y1, x2, y2, fill) in boxes.items():
        d.rounded_rectangle((x1, y1, x2, y2), radius=22, fill=fill, outline="#365F91", width=4)
        bb = d.textbbox((0, 0), name, font=h)
        d.text(((x1+x2-bb[2])/2, (y1+y2-bb[3])/2-4), name, font=h, fill="#17365D")

    def arrow(points, label="", color="#365F91", dashed=False):
        for i in range(len(points)-1):
            d.line([points[i], points[i+1]], fill=color, width=4)
        x1, y1 = points[-2]; x2, y2 = points[-1]
        import math
        a = math.atan2(y2-y1, x2-x1)
        p1 = (x2-18*math.cos(a-0.5), y2-18*math.sin(a-0.5))
        p2 = (x2-18*math.cos(a+0.5), y2-18*math.sin(a+0.5))
        d.polygon([(x2,y2), p1, p2], fill=color)
        if label:
            mx = sum(p[0] for p in points)//len(points)
            my = sum(p[1] for p in points)//len(points)
            d.rounded_rectangle((mx-8, my-18, mx+d.textbbox((0,0),label,font=tiny)[2]+10, my+12), radius=5, fill="white")
            d.text((mx, my-16), label, font=tiny, fill=color)

    arrow([(350,850),(470,850),(470,500),(390,500)], "取指地址")
    arrow([(390,470),(850,300)], "opcode/funct")
    arrow([(390,520),(520,520),(520,610),(600,610)], "rs1/rs2/rd")
    arrow([(390,540),(520,540),(520,980),(600,980)], "imm[31:7]")
    arrow([(900,620),(1120,620)], "rD1/rD2")
    arrow([(900,980),(1020,980),(1020,700),(1120,700)], "ext")
    arrow([(1500,620),(1630,620)], "addr/result")
    arrow([(900,690),(1540,690),(1540,660),(1630,660)], "store data")
    arrow([(1920,610),(2050,610)], "ren/wen/addr")
    arrow([(2200,720),(2200,900)], "read data")
    arrow([(2050,990),(1830,990),(1830,1210),(1570,1210)], "MEXT.ext")
    arrow([(1310,760),(1310,1120)], "ALU.c")
    arrow([(900,1020),(1080,1020),(1080,1250),(1260,1250)], "LUI.ext")
    arrow([(430,1200),(900,1200),(900,1170),(1260,1170)], "PC+4")
    arrow([(1260,1210),(1020,1210),(1020,740),(900,740)], "wD")
    arrow([(120,1200),(70,1200),(70,850),(100,850)], "npc")
    arrow([(350,850),(480,850),(480,1170),(430,1170)], "pc")
    arrow([(600,1010),(500,1010),(500,1240),(430,1240)], "offset")
    arrow([(1120,680),(1020,680),(1020,1320),(470,1320),(470,1270),(430,1270)], "br / jalr target")
    arrow([(350,810),(1040,810),(1040,570),(1120,570)], "PC(AUIPC)")
    arrow([(1050,390),(1050,470),(750,470),(750,520)], "rf/sext/alu control", "#C55A11")
    arrow([(1250,300),(1550,300),(1550,560),(1630,560)], "memory control", "#C55A11")
    arrow([(850,280),(480,280),(480,1150),(430,1150)], "npc_op", "#C55A11")

    d.text((80, 1390), "关键多路选择：ALU-A=RS1/PC；ALU-B=RS2/EXT；WB=ALU/RAM/PC+4/EXT", font=sm, fill="#17365D")
    im.save(DOCS / "miniRV完整数据通路图.png", quality=95)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=55, start=70, bottom=55, end=70):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v)); node.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    rep = OxmlElement("w:tblHeader"); rep.set(qn("w:val"), "true"); tr_pr.append(rep)


def set_run_font(run, size=11, bold=False, color=None, name="Microsoft YaHei"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, 9, color="6B7785")
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    set_run_font(run, 9, color="6B7785")


def configure_section(section, landscape=False):
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width = Inches(11 if landscape else 8.5)
    section.page_height = Inches(8.5 if landscape else 11)
    section.top_margin = Inches(0.65 if landscape else 1.0)
    section.bottom_margin = Inches(0.65 if landscape else 1.0)
    section.left_margin = Inches(0.55 if landscape else 1.0)
    section.right_margin = Inches(0.55 if landscape else 1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    add_page_field(p)


def style_doc(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for style_name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        st = doc.styles[style_name]
        st.font.name = "Calibri"; st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before); st.paragraph_format.space_after = Pt(after)


def add_table(doc, headers, data, widths=None, font_size=8, group_col=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    if widths is None:
        widths = [1/len(headers)] * len(headers)
    total = sum(widths)
    usable = 9.75 if doc.sections[-1].orientation == WD_ORIENT.LANDSCAPE else 6.5
    widths = [usable * x / total for x in widths]
    for j, htxt in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.width = Inches(widths[j]); set_cell_shading(cell, "E8EEF5"); set_cell_margins(cell)
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(htxt); set_run_font(r, font_size, True, BLUE)
    set_repeat_header(table.rows[0])
    for row in data:
        cells = table.add_row().cells
        fill = LIGHT_GREEN if row[0] == "示例" else LIGHT_BLUE if row[0] == "A" else LIGHT_ORANGE
        for j, value in enumerate(row):
            cells[j].width = Inches(widths[j]); cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(cells[j])
            if group_col is not None and j == group_col:
                set_cell_shading(cells[j], fill)
            p = cells[j].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value)); set_run_font(r, font_size, j == 1)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_para(doc, text, bold=False, size=11, color=None, align=None, after=6):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(after); p.paragraph_format.line_spacing = 1.25
    if align is not None: p.alignment = align
    r = p.add_run(text); set_run_font(r, size, bold, color)
    return p


def build_doc():
    doc = Document()
    style_doc(doc)
    configure_section(doc.sections[0], False)

    for _ in range(5): doc.add_paragraph()
    add_para(doc, "计算机设计与实践 · 实验一", True, 13, "7F8C99", WD_ALIGN_PARAGRAPH.CENTER, 18)
    add_para(doc, "支持 miniRV 指令集的\n单周期 CPU 设计与实现", True, 28, "17365D", WD_ALIGN_PARAGRAPH.CENTER, 12)
    add_para(doc, "A、B 两组及 8 条示例指令 · 完整 44 条指令", True, 14, "2E74B5", WD_ALIGN_PARAGRAPH.CENTER, 42)
    add_para(doc, "目标器件：xc7a35tcsg324-1    开发环境：Vivado 2023.2", False, 11, "5A6670", WD_ALIGN_PARAGRAPH.CENTER, 4)
    add_para(doc, "实现内容：数据通路、控制器、访存、乘除法、跳转与自检 Trace", False, 11, "5A6670", WD_ALIGN_PARAGRAPH.CENTER, 4)
    doc.add_page_break()

    doc.add_heading("1. 实验目的与完成范围", level=1)
    add_para(doc, "本设计在课程模板工程基础上完成 miniRV 44 条指令，包含 8 条示例指令、A 组 18 条和 B 组 18 条。CPU 采用模板给定的取指请求/响应、数据访问请求/响应接口；普通算术逻辑和跳转指令在一次执行节拍内完成，存储器及迭代乘除法通过 valid/busy 握手等待完成。")
    for text in ["理解单周期 CPU 的数据通路与控制逻辑。", "掌握 R、I、S、B、U、J 六种指令格式与立即数扩展。", "完成 44 条 miniRV 指令的译码、执行、访存、写回及控制转移。", "使用自建 PC Trace、寄存器写回 Trace 和存储结果进行覆盖验证。"]:
        p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(4); set_run_font(p.add_run(text), 11)

    doc.add_heading("2. 指令范围", level=1)
    add_table(doc, ["分组", "数量", "指令"], [
        ["示例", "8", "SLLI、ADDI、LUI、ORI、LW、JAL、BEQ、BNE"],
        ["A", "18", "SLL、SRL、SRLI、SRA、SRAI、ADD、SUB、AUIPC、XOR、XORI、LB、LBU、LH、LHU、SW、SB、SH、JALR"],
        ["B", "18", "MUL、MULH、MULHU、DIV、DIVU、REM、REMU、SLT、SLTI、SLTU、SLTIU、OR、AND、ANDI、BLT、BGE、BLTU、BGEU"],
    ], [0.8, 0.6, 5.1], 8.5, 0)

    doc.add_heading("3. 完整数据通路", level=1)
    add_para(doc, "完整数据通路以 PC、NPC、RF、SEXT、ALU/MULDIV、MREQ、Data RAM、MEXT 和写回多路器为核心。Controller 由 opcode、funct3、funct7 产生下一 PC、立即数扩展、ALU、访存和寄存器写回控制信号。")
    doc.add_picture(str(DOCS / "miniRV完整数据通路图.png"), width=Inches(6.5))
    p = doc.paragraphs[-1]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_para(doc, "图 1  miniRV 单周期 CPU 完整数据通路", False, 9, "6B7785", WD_ALIGN_PARAGRAPH.CENTER, 4)

    doc.add_page_break()
    doc.add_heading("4. 关键实现说明", level=1)
    key = [
        ("控制器", "完整译码 44 条指令，产生 npc_op、rf_we、rf_wsel、sext_op、alua_sel、alub_sel、alu_op、ram_rop、ram_wop、is_mul 和 is_div。"),
        ("立即数扩展", "支持 I、S、B、U、J 五种立即数拼接与符号扩展；移位立即数使用低 5 位。"),
        ("ALU", "支持加减、逻辑、移位、有符号/无符号比较、六类分支条件以及乘除余数运算。"),
        ("访存", "MREQ 根据地址低两位产生字节使能和对齐判断；MEXT 完成 LB/LH 的符号扩展及 LBU/LHU 的零扩展。"),
        ("跳转", "JAL 目标为 PC+J 型立即数；JALR 目标为 (rs1+I 型立即数)&0xFFFFFFFE；两者写回 PC+4。"),
        ("乘除法", "采用移位加法乘法器和恢复余数除法器；处理除数为 0、-2^31/-1 溢出以及有符号余数符号规则。"),
    ]
    add_table(doc, ["模块", "实现要点"], [["-", a, b] for a,b in []], [1.4, 5.1], 9)
    table = doc.tables[-1]
    for a, b in key:
        cells = table.add_row().cells
        for j, val in enumerate((a,b)):
            set_cell_margins(cells[j]); cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            r = cells[j].paragraphs[0].add_run(val); set_run_font(r, 9, j==0)

    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(sec, True)
    doc.add_heading("5. A、B 两组与示例指令数据通路表", level=1)
    add_para(doc, "表中列出每条指令经过执行单元时的主要数据来源、ALU 功能、访存行为、写回来源及下一 PC 生成方式。")
    headers = ["组", "指令", "型", "ALU-A", "ALU-B", "ALU操作", "访存", "写回", "下一PC"]
    dp = [datapath_row(x) for x in all_insts]
    add_table(doc, headers, dp, [0.5,0.8,0.4,1.0,1.0,0.8,1.8,1.0,2.25], 7.2, 0)

    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(sec, True)
    doc.add_heading("6. 完整控制信号表（编码与主控制）", level=1)
    main_rows = [[r[0],r[1],r[2],r[3],r[4],r[5],r[6],r[7],r[8],r[9]] for r in rows()]
    add_table(doc, ["组","指令","型","opcode","funct3","funct7","npc_op","rf_we","rf_wsel","sext_op"], main_rows, [0.45,0.7,0.35,0.8,0.55,0.85,0.75,0.5,0.7,0.65], 7.0, 0)

    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(sec, True)
    doc.add_heading("7. 完整控制信号表（执行与访存控制）", level=1)
    exe_rows = [[r[0],r[1],r[10],r[11],r[12],r[13],r[14],r[15],r[16]] for r in rows()]
    add_table(doc, ["组","指令","alua_sel","alub_sel","alu_op","ram_rop","ram_wop","is_mul","is_div"], exe_rows, [0.45,0.8,0.8,0.8,0.85,0.75,0.75,0.65,0.65], 7.2, 0)

    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(sec, False)
    doc.add_heading("8. 自检验证", level=1)
    add_para(doc, "由于当前没有课程官方 Basic Trace 包，本工程附带 cpu_core_all_tb.v 和自动生成脚本 gen_all_test.py。测试程序包含全部 44 条指令，参考模型生成期望 PC Trace 与寄存器写回 Trace，并额外检查 SW、SB、SH 的最终存储结果。")
    add_table(doc, ["检查项", "覆盖结果"], [
        ["指令助记符", "44/44"],
        ["自检程序机器指令", "81 条"],
        ["有效 PC Trace", "73 项（已考虑跳转和已取分支）"],
        ["寄存器写回 Trace", "63 项"],
        ["存储检查", "SW、SB、SH 三种写掩码及数据对齐"],
        ["乘除法算法检查", "随机/边界输入检查通过"],
        ["完整 RTL 功能仿真", "44 条指令自检通过"],
    ], [2.1,4.4], 9)
    add_para(doc, "Vivado 运行：在工程目录的 Tcl Console 执行 source run_all_tb.tcl。若仿真无误，控制台输出 ALL 44 miniRV INSTRUCTIONS PASSED。官方验收前仍需使用学校提供的 Basic Trace 再次验证接口与 Trace 规范。", True, 10, "17365D")

    doc.add_page_break()
    doc.add_heading("9. 工程文件说明", level=1)
    add_table(doc, ["文件", "作用"], [
        ["src/rtl/Controller.v", "44 条指令译码与控制信号生成"],
        ["src/rtl/ALU.v", "算术、逻辑、移位、比较、分支和乘除法结果选择"],
        ["src/rtl/SEXT.v", "I/S/B/U/J 立即数扩展"],
        ["src/rtl/MREQ.v、MEXT.v", "访存掩码、对齐及读数据扩展"],
        ["src/rtl/multiplier.v、divider.v", "迭代乘法器、无符号除法器"],
        ["src/sim/cpu_core_all_tb.v", "44 条指令自检 testbench"],
        ["src/coe/all_instructions_test.*", "自检汇编、HEX 与 COE"],
        ["docs/*.csv、*.png", "控制表、数据通路表与完整数据通路图"],
    ], [2.5,4.0], 9)

    props = doc.core_properties
    props.title = "实验一：支持 miniRV 44 条指令的单周期 CPU 设计与实现"
    props.subject = "计算机设计与实践"
    props.author = ""
    props.keywords = "miniRV, RISC-V, 单周期CPU, Vivado, Verilog"
    doc.save(OUT)


def write_readme():
    text = """# miniRV 44 条指令完整工程

本工程基于课程 `miniRV_basic` 模板补全 44 条 miniRV 指令，包含 8 条示例指令、A 组 18 条和 B 组 18 条。

## 已完成

- `Controller.v`：44 条指令译码与全部控制信号
- `ALU.v`：算术、逻辑、移位、比较、分支、乘除余数
- `SEXT.v`：I/S/B/U/J 型立即数
- `MREQ.v`、`MEXT.v`：LB/LBU/LH/LHU/LW/SB/SH/SW
- `NPC.v`：顺序、分支、JAL、JALR
- `multiplier.v`、`divider.v`：迭代乘除法
- `cpu_core_all_tb.v`：PC Trace、寄存器写回 Trace 与存储器自检

## 自检运行

1. 使用 Vivado 2023.2 打开 `miniRV.xpr`。
2. 在 Tcl Console 中执行：`source run_all_tb.tcl`。
3. 通过时控制台输出：`ALL 44 miniRV INSTRUCTIONS PASSED`。

本交付版本已使用 Icarus Verilog 编译并运行上述自检，结果为全部 44 条指令通过。

如修改测试程序，执行 `python tools/gen_all_test.py` 可重新生成汇编、HEX、COE 和 testbench。

## 设计资料

- `docs/miniRV_44条指令控制信号表.csv`
- `docs/miniRV_44条指令数据通路表.csv`
- `docs/miniRV完整数据通路图.png`
- `docs/self_test_coverage.txt`

当前自检不是课程官方 Basic Trace。取得官方测试包后，应以官方 Trace 结果作为最终验收依据。
"""
    (ROOT / "README.md").write_text(text, encoding="utf-8")


if __name__ == "__main__":
    write_csvs()
    draw_diagram()
    write_readme()
    build_doc()
